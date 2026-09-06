#!/usr/bin/env python3
"""
One-off generator for test/fixtures/task_a_dataset.json — a plain-JSON dump of
Task_A_Raw_Placement_Data_5_Years.xlsx shaped exactly like what index.html's readFile()
callback hands to the rest of the pipeline (sheetNames, sheetHeaders, and a flat row list
where each row is tagged with __sheet). Committing this JSON lets test/regression.test.js
run without any xlsx-parsing dependency at test time — this script only needs to be
re-run if the source workbook itself changes.

Usage: python3 test/fixtures/generate_fixture.py
Requires: openpyxl (not needed to run the tests themselves, only to regenerate this fixture).
"""
import json
import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, 'Task_A_Raw_Placement_Data_5_Years.xlsx')
OUT = os.path.join(ROOT, 'test', 'fixtures', 'task_a_dataset.json')

wb = openpyxl.load_workbook(SRC, data_only=True)
sheetNames = wb.sheetnames
sheetHeaders = []
allRows = []
for name in sheetNames:
    ws = wb[name]
    rows_iter = ws.iter_rows(values_only=True)
    header = [h if h is not None else '' for h in next(rows_iter)]
    sheetHeaders.append(header)
    for row in rows_iter:
        if row is None or all(v is None for v in row):
            continue
        obj = {h: ('' if v is None else v) for h, v in zip(header, row)}
        obj['__sheet'] = name
        allRows.append(obj)

with open(OUT, 'w') as f:
    json.dump({'sheetNames': sheetNames, 'sheetHeaders': sheetHeaders, 'rows': allRows}, f, default=str)

print(f'Wrote {len(allRows)} rows across {len(sheetNames)} sheets to {OUT}')
