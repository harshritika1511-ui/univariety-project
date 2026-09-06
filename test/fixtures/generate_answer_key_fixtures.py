#!/usr/bin/env python3
"""
One-off generator for the JSON dumps of the two multi-sheet XLSX answer-key fixtures
(files/Placement_TestCase_MultiSheet_YearTabs.xlsx, files/Placement_TestCase_MultiSheet_Trap.xlsx),
shaped the same way as test/fixtures/generate_fixture.py — sheetNames/sheetHeaders/rows with
each row tagged __sheet — so test/answer-key.test.js can run without an xlsx-parsing
dependency at test time. Re-run only if those source files change.

Usage: python3 test/fixtures/generate_answer_key_fixtures.py
Requires: openpyxl (not needed to run the tests themselves, only to regenerate these fixtures).
"""
import json
import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

SOURCES = [
    ('files/Placement_TestCase_MultiSheet_YearTabs.xlsx', 'test/fixtures/multisheet_yeartabs.json'),
    ('files/Placement_TestCase_MultiSheet_Trap.xlsx', 'test/fixtures/multisheet_trap.json'),
]

for src_rel, out_rel in SOURCES:
    src = os.path.join(ROOT, src_rel)
    out = os.path.join(ROOT, out_rel)
    wb = openpyxl.load_workbook(src, data_only=True)
    sheetNames = wb.sheetnames
    sheetHeaders = []
    allRows = []
    for name in sheetNames:
        ws = wb[name]
        rows_iter = ws.iter_rows(values_only=True)
        header = [h if h is not None else '' for h in next(rows_iter)]
        sheetHeaders.append(header)
        for row in rows_iter:
            if row is None or all(v is None for v in row):
                continue
            obj = {h: ('' if v is None else v) for h, v in zip(header, row)}
            obj['__sheet'] = name
            allRows.append(obj)
    with open(out, 'w') as f:
        json.dump({'sheetNames': sheetNames, 'sheetHeaders': sheetHeaders, 'rows': allRows}, f, default=str)
    print(f'Wrote {len(allRows)} rows across {len(sheetNames)} sheets: {src_rel} -> {out_rel}')
